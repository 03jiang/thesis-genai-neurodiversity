from pathlib import Path

DATA_DIR = Path(__file__).resolve().parents[2] / "data"
CN_PATH = DATA_DIR / "chinese_survey.xlsx"
EN_PATH = DATA_DIR / "english_survey.xlsx"
OUT_PATH = DATA_DIR / "combined_75.csv"

from openpyxl import load_workbook
import pandas as pd


def parse_cn_single(cell):
    if not cell or "跳过" in str(cell):
        return None
    return str(cell).strip()[0]


def parse_cn_multi(cell):
    if not cell or "跳过" in str(cell):
        return set()
    return {part.strip()[0] for part in str(cell).split("┋") if part.strip()}


EN_LETTER = {1.0: "A", 2.0: "B", 3.0: "C", 4.0: "D", 5.0: "E"}
EN_Q5     = {1.0: "Yes", 2.0: "No", 3.0: "Prefer"}


def load_cn(path):
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]

    q5_char_to_label = {"A": "Yes", "B": "No", "C": "Prefer"}
    records = []
    for i, row in enumerate(data, start=1):
        if not row[6] or not str(row[6]).startswith("A"):
            continue
        q5_char = parse_cn_single(row[11])
        records.append({
            "respondent_id": f"CN_{i:03d}",
            "source":        "CN",
            "Q1":            parse_cn_single(row[7]),
            "Q2":            parse_cn_single(row[8]),
            "Q3":            parse_cn_single(row[9]),
            "Q4_set":        parse_cn_multi(row[10]),
            "Q5_label":      q5_char_to_label.get(q5_char),
        })
    return records


def load_en(path):
    wb = load_workbook(path, read_only=True)
    ws = wb["Dataset"]
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]

    records = []
    for i, row in enumerate(data, start=1):
        if row[13] != 1.0:
            continue
        q4_set = {letter for letter, val in zip("ABCDEF", row[3:9])
                  if val == 1.0}
        records.append({
            "respondent_id": f"EN_{i:03d}",
            "source":        "EN",
            "Q1":            EN_LETTER.get(row[0]),
            "Q2":            EN_LETTER.get(row[1]),
            "Q3":            EN_LETTER.get(row[2]),
            "Q4_set":        q4_set,
            "Q5_label":      EN_Q5.get(row[9]),
        })
    return records


def build_dataframe(cn_path, en_path):
    """合并两份数据，做一热编码，返回最终 DataFrame。"""
    records = load_cn(cn_path) + load_en(en_path)
    df = pd.DataFrame(records)

    for q, letters in [("Q1", "ABCDE"), ("Q2", "ABCD"), ("Q3", "ABC")]:
        for L in letters:
            df[f"{q}_{L}"] = (df[q] == L).astype(int)

    for L in "ABCDEF":
        df[f"Q4_{L}"] = df["Q4_set"].apply(lambda s: int(L in s))

    df = df.drop(columns=["Q1", "Q2", "Q3", "Q4_set"])
    cols = (["respondent_id", "source"]
            + [f"Q1_{L}" for L in "ABCDE"]
            + [f"Q2_{L}" for L in "ABCD"]
            + [f"Q3_{L}" for L in "ABC"]
            + [f"Q4_{L}" for L in "ABCDEF"]
            + ["Q5_label"])
    return df[cols]



def sanity_check(df):
    q1_sum = df[[f"Q1_{L}" for L in "ABCDE"]].sum(axis=1)
    q2_sum = df[[f"Q2_{L}" for L in "ABCD"]].sum(axis=1)
    q3_sum = df[[f"Q3_{L}" for L in "ABC"]].sum(axis=1)
    q4_sum = df[[f"Q4_{L}" for L in "ABCDEF"]].sum(axis=1)

    assert (q1_sum == 1).all(), f"Q1 应单选，但存在异常行: {df[q1_sum != 1]}"
    assert (q2_sum == 1).all(), f"Q2 应单选，但存在异常行: {df[q2_sum != 1]}"
    assert (q3_sum == 1).all(), f"Q3 应单选，但存在异常行: {df[q3_sum != 1]}"

    q4_over = df[q4_sum > 3]
    if len(q4_over) > 0:
        print(f" {len(q4_over)} 位受访者 Q4 选了 >3 项（平台未强制上限）")
        print(f"   受影响 ID: {list(q4_over['respondent_id'])}")
    print(f"一致性校验通过")


def compare_with_report(df):
    """打印频次表，对照 PDF 报告。"""
    print("\n" + "=" * 60)
    print("与 PDF 报告 (survey_combined_report_en.pdf) 频次比对")
    print("=" * 60)
    print(f"总样本 n = {len(df)}   （报告 n=76；差 1 是那位英文 Partial）")
    print(f"来源分布: {dict(df['source'].value_counts())}")

    def show(label, cols):
        counts = {c.split("_")[1]: int(df[c].sum()) for c in cols}
        print(f"\n{label}: {counts}")

    show("Q1 (报告合并: A=25 B=19 C=12 D=8 E=12 — 那位 Partial 选了 A，我们少 1)",
         [f"Q1_{L}" for L in "ABCDE"])
    show("Q2 (报告合并: A=44 B=19 C=5 D=8 — Partial 选了 A，我们少 1)",
         [f"Q2_{L}" for L in "ABCD"])
    show("Q3 (报告合并: A=43 B=10 C=23 — Partial 选了 A，我们少 1)",
         [f"Q3_{L}" for L in "ABC"])
    show("Q4 (报告合并: A=56 B=24 C=11 D=15 E=41 F=38 — Q4 报告本就是 n=75)",
         [f"Q4_{L}" for L in "ABCDEF"])

    print(f"\nQ5 分布: {dict(df['Q5_label'].value_counts(dropna=False))}")


if __name__ == "__main__":
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    df = build_dataframe(CN_PATH, EN_PATH)
    sanity_check(df)
    compare_with_report(df)

    df.to_csv(OUT_PATH, index=False)
    print(f"\n已保存 → {OUT_PATH}   （{len(df)} 行 × {len(df.columns)} 列）")
    print(f"\n前 3 行预览：")
    print(df.head(3).to_string())
